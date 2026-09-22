"""Génère la copie Studi éditable à partir de l'état réellement vérifié."""

import os
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
NOM_FICHIER = os.getenv("STUDI_NOM", "NOM").upper()
PRENOM_FICHIER = os.getenv("STUDI_PRENOM", "Prenom")
OUTPUT = ROOT / "output" / f"ECF_TPDeveloppeurWebEtWebMobile_copiearendre_{NOM_FICHIER}_{PRENOM_FICHIER}.docx"


def add_question(doc, title, answer):
    doc.add_heading(title, level=2)
    doc.add_paragraph(answer)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Arial"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(20)
styles["Heading 1"].font.size = Pt(14)
styles["Heading 2"].font.size = Pt(11)
for border in styles["Title"]._element.findall(".//" + qn("w:pBdr")):
    border.getparent().remove(border)

doc.add_paragraph("Copie à rendre Studi Vite et Gourmand", style="Title")
if all(os.getenv(name) for name in ("STUDI_NOM", "STUDI_PRENOM", "STUDI_NAISSANCE")):
    doc.add_paragraph(
        f"Nom : {os.environ['STUDI_NOM']}    Prénom : {os.environ['STUDI_PRENOM']}    "
        f"Date de naissance : {os.environ['STUDI_NAISSANCE']}"
    )
doc.add_paragraph(
    "TP Développeur Web et Web Mobile. Cette copie décrit l'état vérifié du projet au "
    "22 septembre 2026. Les éléments non encore publiés ou non validés sont indiqués comme tels."
)

doc.add_heading("Liens et accès demandés", level=1)
for label, value in [
    ("Dépôt GitHub public, version jury", "https://github.com/mehdi-prieux/vite-et-gourmand/tree/livraison-studi"),
    ("Outil de gestion de projet", "https://github.com/mehdi-prieux/vite-et-gourmand/issues — suivi public reconstitué au 22 septembre 2026"),
    ("Application publique", "NON DÉPLOYÉE — recette locale uniquement à ce stade"),
    ("Administrateur de démonstration", "admin@test.com / Admin!Demo2026 — compte créé par le SQL d'insertion, réservé à la démonstration"),
]:
    p = doc.add_paragraph()
    p.add_run(label + " : ").bold = True
    p.add_run(value)

doc.add_heading("Partie 1 Analyse des besoins", level=1)
add_question(doc, "1 Résumé du projet", """Vite et Gourmand est une application web destinée à un traiteur bordelais. Elle permet à un visiteur de découvrir les menus, leurs compositions, allergènes, prix, disponibilités et avis validés. Le visiteur peut contacter l'entreprise ou créer un compte. Le client connecté renseigne les informations nécessaires à sa commande, choisit une date de prestation, consulte le prix calculé et retrouve ensuite son historique. Il peut modifier ou annuler une demande avant son acceptation et déposer un avis après une commande terminée. Le personnel dispose d'un espace pour suivre les commandes, faire évoluer leur statut, gérer le catalogue, les horaires et modérer les avis. Un administrateur peut également gérer les comptes employés et consulter des statistiques filtrées.

L'application repose sur une interface HTML, CSS et JavaScript, une API PHP et une base MySQL pour les opérations métier. Une projection documentaire minimale vers CouchDB est préparée pour les statistiques ; sa validation de bout en bout sur une instance réelle reste à effectuer. Les mots de passe sont hachés, les accès contrôlés par rôle et les écritures authentifiées protégées contre les requêtes intersites. Les parcours principaux ont été testés localement sur une base dédiée. Les e-mails ont été vérifiés en mode journal de test, pas encore avec un transport réel. Le projet n'est pas présenté comme déployé tant qu'une URL publique et une recette hébergée ne sont pas disponibles.""")
add_question(doc, "2 Cahier des charges et besoin", """Le service doit présenter le catalogue et les informations de contact, permettre l'inscription et la connexion, enregistrer une commande avec règles de prix et de livraison, puis afficher son suivi. Le client peut gérer ses demandes autorisées et publier un avis soumis à modération. Le personnel gère les commandes, menus, plats, allergènes, horaires et avis. L'administrateur gère les employés et les statistiques. Les données transactionnelles restent dans MySQL ; les indicateurs de commandes et de chiffre d'affaires doivent être lus depuis une base non relationnelle distincte. Les exigences transversales sont l'affichage responsive, l'accessibilité, la sécurité des accès, les mentions légales et CGV, les e-mails métier, la documentation et un déploiement accessible au jury. Les détails de tests et limites sont dans docs/matrice-recette-finale-studi.md.""")

heading = doc.add_heading("Partie 2 Spécifications techniques", level=1)
heading.paragraph_format.page_break_before = True
add_question(doc, "1 Technologies utilisées et justification", """HTML, CSS et JavaScript natifs fournissent une interface responsive sans chaîne de compilation. PHP 8.1 ou supérieur expose l'API métier ; PDO et les requêtes préparées accèdent à MySQL, choisi pour l'intégrité des utilisateurs, commandes, menus et stocks. CouchDB a été choisi comme base documentaire non relationnelle : son API HTTP est accessible depuis PHP avec cURL et permet de stocker uniquement les champs utiles aux statistiques. Le code de projection et de lecture est présent, mais le résultat réel sur CouchDB n'est pas encore validé. L'envoi réel utilise actuellement la fonction mail() de PHP ; la recette locale utilise un journal, et un transport fiable reste à configurer chez l'hébergeur.""")
add_question(doc, "2 Environnement de travail", """Le dépôt Git contient frontend/, backend/, database/mysql/, docs/, tests/ et output/pdf/. Le serveur local se lance depuis la racine avec php -S 127.0.0.1:8000. Les connexions utilisent des variables d'environnement ou un fichier local ignoré par Git ; .env.example documente les paramètres. Le README détaille l'installation et les scripts SQL de création et d'insertion. Les tests avec écriture vérifient que SELECT DATABASE() renvoie vite_gourmand_test avant toute mutation. Les branches main, develop, fix/database-config-env et livraison-studi existent réellement sur le dépôt ; aucun historique de branche n'a été inventé.""")
add_question(doc, "3 Mécanismes de sécurité", """Les mots de passe sont hachés avec password_hash() et vérifiés avec password_verify(). Les sessions serveur et les contrôles de rôle limitent les API client, employé et administrateur. Les opérations authentifiées vérifient un jeton CSRF. Les entrées sont validées côté serveur, les requêtes SQL sont préparées et les commandes utilisent des transactions. Les prix et frais sont recalculés côté serveur. Les avis restent invisibles au public avant validation. Les jetons de réinitialisation sont à usage limité. HTTPS est exigé en production. Les comptes de démonstration ne doivent pas servir à une exploitation réelle.""")
add_question(doc, "4 Veille sur les vulnérabilités", """La veille de sécurité a été orientée vers les risques concrets d'une application PHP : injections SQL, accès non autorisés, CSRF, fuite de secrets et configuration de session. Les recommandations OWASP Top 10 et PHP Security ont guidé les contrôles effectués dans l'API et les tests de rôles. Sources : https://owasp.org/www-project-top-ten/ et https://www.php.net/manual/en/security.php. Cette veille ne constitue pas un audit de sécurité externe.""")

heading = doc.add_heading("Partie 3 Recherche", level=1)
heading.paragraph_format.page_break_before = True
add_question(doc, "1 Situation de recherche sur un site anglophone", """La synchronisation des statistiques oblige à mettre à jour plusieurs fois le même document de commande dans CouchDB après des changements de statut. La documentation officielle anglophone CouchDB a été consultée pour comprendre le fonctionnement des révisions et éviter les conflits lors d'une mise à jour. Source : Apache CouchDB, « /{db}/{docid} », https://docs.couchdb.org/en/stable/api/document/common.html. Cette recherche a conduit à relire la révision courante avant une mise à jour ; le fonctionnement réel reste à contrôler sur une instance CouchDB.""")
add_question(doc, "2 Extrait et traduction", """Extrait : « To update an existing document you must specify the current revision number within the _rev parameter. » Traduction : « Pour mettre à jour un document existant, il faut indiquer le numéro de révision actuel dans le paramètre _rev. » Source : documentation officielle Apache CouchDB, section « Updating an Existing Document », https://docs.couchdb.org/en/stable/api/document/common.html.""")

doc.add_heading("Partie 4 Informations complémentaires", level=1)
add_question(doc, "1 Autres ressources", """README.md ; database/mysql/create_database.sql ; database/mysql/insert_data.sql ; output/pdf/manuel_utilisateur_studi.pdf ; output/pdf/charte_graphique_et_maquettes_studi.pdf ; docs/documentation-technique-studi.md ; docs/gestion-projet-studi.md ; docs/deploiement-studi.md ; docs/matrice-recette-finale-studi.md.""")
add_question(doc, "2 État de remise", """Les parcours locaux testés et les documents disponibles sont distingués des points non validés : synchronisation CouchDB réelle, transport e-mail réel, identité légale définitive et recette sur URL publique. La branche publique livraison-studi contient les commits de cette version ; main reste inchangée et la PR n°2 n'est pas fusionnée. Aucun résultat de déploiement ou de statistique réelle n'est affirmé sans preuve.""")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Vite et Gourmand — copie Studi éditable")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
